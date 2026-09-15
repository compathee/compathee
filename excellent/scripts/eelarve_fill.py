#!/usr/bin/env python3
"""Fill an Eelarve workbook from monthly Kasumiaruanne source exports.

This is the primary Windows-client flow. It writes object amounts into the
matching konto block and month column. The legacy unpivot combiner is separate.
"""

from __future__ import annotations

import argparse
import csv
import errno
import re
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Callable

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet
except ModuleNotFoundError as exc:  # pragma: no cover - exercised by users without deps
    raise SystemExit(
        "Missing dependency: openpyxl. Install dependencies with "
        "`python -m pip install -r requirements.txt`."
    ) from exc


MONTHS = [
    "Jaanuar",
    "Veebruar",
    "Märts",
    "Aprill",
    "Mai",
    "Juuni",
    "Juuli",
    "August",
    "September",
    "Oktoober",
    "November",
    "Detsember",
]
MONTHS_LOOKUP = {month.casefold(): month for month in MONTHS}
SOURCE_SHEET_NAME = "Kasumiaruanne"
OBJECT_HEADER_RE = re.compile(
    r"^[A-Z0-9ÕÄÖÜŠŽ]+[_-][A-Z0-9ÕÄÖÜŠŽõäöüšž_-]+",
    re.IGNORECASE,
)
SKIP_SOURCE_HEADERS = {
    "konto",
    "account",
    "nimetus",
    "description",
    "name",
    "periood",
    "period",
    "kokku",
    "total",
    "objekt",
    "object",
    "summa",
}
YEAR_RE = re.compile(r"^(?:19|20)\d{2}$")
KONTO_RE = re.compile(r"\d{3,6}")
SUM_CALL_RE = re.compile(r"SUM\(([^)]+)\)", re.IGNORECASE)
RANGE_RE = re.compile(
    r"(\$?[A-Za-z]+\$?\d+)\s*:\s*(\$?[A-Za-z]+\$?\d+)"
)
CELL_RE = re.compile(r"(\$?)([A-Za-z]+)(\$?)(\d+)")


@dataclass(frozen=True)
class SourceFact:
    konto: str
    konto_description: str
    object_code: str
    amount: float


@dataclass(frozen=True)
class FillStats:
    written: int
    rows_added: int
    missing_konto_blocks: int
    output_path: Path


@dataclass(frozen=True)
class EelarveLayout:
    header_row: int
    konto_col: int
    object_col: int
    month_columns: dict[str, int]
    total_col: int | None


@dataclass
class KontoBlock:
    konto: str
    header_row: int
    end_row: int
    total_rows: list[int]
    labels: dict[str, int]


def project_root() -> Path:
    return Path(__file__).resolve().parent.parent


def ensure_work_dirs(root: Path | None = None) -> tuple[Path, Path]:
    base = project_root() if root is None else root
    input_dir = base / "input"
    output_dir = base / "output"
    input_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    return input_dir, output_dir


def default_new_eelarve_path(root: Path | None = None, year: int | None = None) -> Path:
    base = project_root() if root is None else root
    resolved_year = date.today().year if year is None else year
    return base / "output" / f"Eelarve_{resolved_year}.xlsx"


class JobAborted(Exception):
    """Raised when the user aborts after a locked/busy file prompt or Ctrl+C."""


INTERRUPT_MESSAGE = "Работа прервана (Ctrl+C).\nJob interrupted."
CANCEL_CHECK_EVERY = 25


def progress_percent(completed: int, total: int, *, start: int = 30, end: int = 90) -> int:
    if total <= 0:
        return end
    completed = min(max(int(completed), 0), int(total))
    return start + (end - start) * completed // total


def is_lock_error(exc: BaseException) -> bool:
    if isinstance(exc, PermissionError):
        return True
    if not isinstance(exc, OSError):
        return False
    if getattr(exc, "errno", None) in {errno.EACCES, errno.EPERM, errno.EBUSY}:
        return True
    if getattr(exc, "winerror", None) in {5, 32, 33}:
        return True
    text = str(exc).casefold()
    return any(
        needle in text
        for needle in (
            "being used by another",
            "used by another process",
            "permission denied",
            "file is locked",
            "заблокирован",
            "занят",
        )
    )


class ProgressReporter:
    def __init__(self, enabled: bool = True, stream: Any | None = None) -> None:
        self.enabled = enabled
        self.stream = stream or sys.stdout
        self.percent = -1
        self._fill_done = 0
        self._fill_total = 0

    def stage(self, percent: int, message: str) -> None:
        if not self.enabled:
            return
        percent = max(0, min(100, int(percent)))
        if percent < self.percent:
            percent = self.percent
        self.percent = percent
        print(f"[{self.percent:3d}%] {message}", file=self.stream, flush=True)

    def begin_fill(self, total: int) -> None:
        self._fill_total = max(0, int(total))
        self._fill_done = 0
        self.stage(30 if self._fill_total else 90, f"Начато заполнение: {self._fill_total} сумм")

    def add_filled(self, n: int = 1) -> None:
        self._fill_done += n
        total = self._fill_total
        if total <= 0:
            return
        step = max(1, min(25, total // 20 or 1))
        if self._fill_done in {1, total} or self._fill_done % step == 0:
            self.stage(
                progress_percent(self._fill_done, total, start=30, end=90),
                f"Заполнение Eelarve… ({self._fill_done} из {total})",
            )


def _resolved_path_text(path: Path) -> str:
    try:
        return str(Path(path).expanduser().resolve())
    except OSError:
        return str(path)


def ask_lock_action(path: Path, *, use_gui: bool = False, kind: str = "source") -> str:
    resolved = _resolved_path_text(path)
    message = (
        f"Файл занят (открыт в Excel или заблокирован):\n{resolved}\n"
        f"File is busy/locked:\n{resolved}"
    )
    print(message, flush=True)
    retry_kinds = kind in {"output", "eelarve"}
    if retry_kinds:
        print("1 = повторить / продолжить,  2 = прервать  (y/n)", flush=True)
    else:
        print("1 = продолжить (пропустить файл),  2 = прервать  (y/n)", flush=True)

    if use_gui:
        try:
            import tkinter as tk
            from tkinter import messagebox

            root = tk.Tk()
            root.withdraw()
            try:
                extra = "\n\nПовторить?" if retry_kinds else "\n\nПропустить этот файл и продолжить?"
                if retry_kinds:
                    return "continue" if messagebox.askretrycancel("Файл занят", message + extra) else "abort"
                return "continue" if messagebox.askyesno("Файл занят", message + extra) else "abort"
            finally:
                root.destroy()
        except Exception:
            pass

    try:
        answer = input("Выбор / choice: ").strip().casefold()
    except EOFError:
        return "abort"
    if answer in {"1", "y", "yes", "c", "continue", "retry", "r"}:
        return "continue"
    return "abort"


def run_with_lock_handling(
    operation: Callable[[], Any],
    path: Path,
    *,
    kind: str,
    handle_locks: bool,
    use_gui: bool,
) -> tuple[str, Any]:
    if not handle_locks:
        return "ok", operation()
    while True:
        try:
            return "ok", operation()
        except Exception as exc:
            if not is_lock_error(exc):
                raise
            action = ask_lock_action(path, use_gui=use_gui, kind=kind)
            if action == "abort":
                raise JobAborted(
                    "Работа прервана: файл занят.\n"
                    f"Job aborted: file is locked.\n{_resolved_path_text(path)}"
                ) from exc
            if kind in {"source", "objects"}:
                return "skip", None
            try:
                return "ok", operation()
            except Exception as exc2:
                if not is_lock_error(exc2):
                    raise
                continue


def notify_user(
    title: str,
    message: str,
    *,
    error: bool = False,
    use_gui: bool = False,
) -> None:
    stream = sys.stderr if error else sys.stdout
    print(message, file=stream)
    if not use_gui:
        return
    try:
        import tkinter as tk
        from tkinter import messagebox

        root = tk.Tk()
        root.withdraw()
        try:
            if error:
                messagebox.showerror(title, message)
            else:
                messagebox.showinfo(title, message)
        finally:
            root.destroy()
    except Exception:
        return


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        try:
            as_number = float(text)
            if as_number.is_integer():
                return str(int(as_number))
        except ValueError:
            pass
    return text


def normalize_header(value: Any, fallback_index: int) -> str:
    text = cell_text(value)
    if not text:
        text = f"column_{fallback_index}"
    return re.sub(r"\s+", " ", text)


def parse_amount(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\u00a0", " ").replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_konto(value: Any) -> str | None:
    text = cell_text(value)
    if not text:
        return None
    for match in KONTO_RE.findall(text):
        if YEAR_RE.fullmatch(match):
            continue
        return match
    return None


def is_total_label(text: str) -> bool:
    normalized = re.sub(r"\s+", " ", (text or "").replace("\u00a0", " ")).strip().casefold()
    return bool(re.match(r"^kokku\b", normalized))


def is_object_header(header: str) -> bool:
    clean = header.strip()
    if not clean:
        return False
    code = re.sub(r"\s*\(.*\)\s*$", "", clean).strip()
    lowered = code.casefold()
    if lowered in SKIP_SOURCE_HEADERS or lowered.startswith("kokku"):
        return False
    return bool(OBJECT_HEADER_RE.match(code))


def parse_object_header(header: str) -> str:
    text = header.strip()
    match = re.match(r"^(?P<code>.+?)\s*\((?P<name>.+)\)\s*$", text)
    if match:
        return match.group("code").strip()
    return re.sub(r"\s*\(.*\)\s*$", "", text).strip()


def parse_month(value: str) -> str:
    text = (value or "").strip()
    if not text:
        raise ValueError("Month is required")
    if text.isdigit():
        index = int(text)
        if 1 <= index <= 12:
            return MONTHS[index - 1]
        raise ValueError(f"Month number out of range: {text}")
    canonical = MONTHS_LOOKUP.get(text.casefold())
    if canonical is None:
        raise ValueError(f"Unknown month: {text}")
    return canonical


def load_object_lookup(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    path = path.expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"Object lookup file not found: {path}")

    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = [cell_text(value) for value in next(iterator)]
        except StopIteration:
            return {}
        rows = [dict(zip(header, row, strict=False)) for row in iterator]
    else:
        raise ValueError("Object lookup must be .csv or .xlsx")

    lookup: dict[str, str] = {}
    for row in rows:
        lowered = {str(key).strip().lower(): value for key, value in row.items()}
        code = cell_text(
            lowered.get("object_code")
            or lowered.get("code")
            or lowered.get("object")
            or lowered.get("objekt")
        )
        name = cell_text(
            lowered.get("object_name")
            or lowered.get("name")
            or lowered.get("description")
            or lowered.get("nimetus")
        )
        if code:
            lookup[code] = name
    return lookup


def _normalize_match_text(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\u00a0", " ")).strip().casefold()


def matches_label(cell: str, label: str) -> bool:
    haystack = _normalize_match_text(cell)
    needle = _normalize_match_text(label)
    if not haystack or not needle:
        return False
    if haystack == needle:
        return True
    pattern = r"(?<![0-9a-zõäöüšž_])" + re.escape(needle) + r"(?![0-9a-zõäöüšž_])"
    return bool(re.search(pattern, haystack))


def find_sheet_by_name(workbook, name: str):
    wanted = name.casefold()
    for worksheet in workbook.worksheets:
        if worksheet.title.casefold() == wanted:
            return worksheet
    raise ValueError(f"Sheet {name!r} not found in {workbook.sheetnames}")


def parse_kasumiaruanne(path: Path) -> list[SourceFact]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = find_sheet_by_name(workbook, SOURCE_SHEET_NAME)
    header_row = 1
    headers: list[str] = []
    best_score = -1
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1
    ):
        candidate = [normalize_header(value, index + 1) for index, value in enumerate(row)]
        score = sum(1 for header in candidate if is_object_header(header))
        if score > best_score:
            best_score = score
            header_row = row_number
            headers = candidate
        if score:
            break
    object_columns = [
        (index, parse_object_header(header))
        for index, header in enumerate(headers)
        if is_object_header(header)
    ]
    if not object_columns:
        raise ValueError(f"No object columns found in {path.name} sheet Kasumiaruanne")

    facts: list[SourceFact] = []
    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        values = list(row)
        col_a = values[0] if values else None
        col_b = values[1] if len(values) > 1 else None
        if is_total_label(cell_text(col_a)) or is_total_label(cell_text(col_b)):
            continue
        konto = parse_konto(col_a) or parse_konto(col_b)
        if not konto:
            continue
        description = cell_text(col_b)
        for index, code in object_columns:
            amount = parse_amount(values[index] if index < len(values) else None)
            if amount is None or amount == 0:
                continue
            facts.append(
                SourceFact(
                    konto=konto,
                    konto_description=description,
                    object_code=code,
                    amount=amount,
                )
            )
    return facts


def _row_values(worksheet: Worksheet, row: int) -> list[Any]:
    return [worksheet.cell(row, column).value for column in range(1, worksheet.max_column + 1)]


def last_used_row(worksheet: Worksheet, layout: EelarveLayout | None = None) -> int:
    if layout is None:
        columns = range(1, min(worksheet.max_column, 16) + 1)
    else:
        columns = {layout.konto_col, layout.object_col}
        columns.update(layout.month_columns.values())
        if layout.total_col:
            columns.add(layout.total_col)
    for row in range(worksheet.max_row, 0, -1):
        if any(cell_text(worksheet.cell(row, column).value) for column in columns):
            return row
    return 1


def detect_eelarve_layout(worksheet: Worksheet) -> EelarveLayout:
    best_row = 1
    best_headers: list[str] = []
    best_score = -1
    scan_to = min(worksheet.max_row, 40) or 1
    for row in range(1, scan_to + 1):
        headers = [
            normalize_header(worksheet.cell(row, column).value, column)
            for column in range(1, worksheet.max_column + 1)
        ]
        score = sum(1 for header in headers if header.casefold() in MONTHS_LOOKUP)
        if score > best_score:
            best_score = score
            best_row = row
            best_headers = headers
    if best_score < 1:
        raise ValueError(f"No month headers found in sheet {worksheet.title!r}")

    month_columns: dict[str, int] = {}
    konto_col = 1
    object_col = 2 if len(best_headers) > 1 else 1
    total_col: int | None = None
    for index, header in enumerate(best_headers, start=1):
        lowered = header.casefold()
        month = MONTHS_LOOKUP.get(lowered)
        if month:
            month_columns[month] = index
            continue
        if lowered in {"konto", "account"}:
            konto_col = index
        elif lowered in {"object", "objekt", "nimetus"}:
            object_col = index
        elif lowered in {"kokku", "total"} or lowered.startswith("kokku") or "aasta" in lowered:
            if total_col is None:
                total_col = index
    return EelarveLayout(
        header_row=best_row,
        konto_col=konto_col,
        object_col=object_col,
        month_columns=month_columns,
        total_col=total_col,
    )


def pick_eelarve_sheet(workbook) -> Worksheet:
    best_sheet = workbook.worksheets[0]
    best_score = -1
    for worksheet in workbook.worksheets:
        try:
            layout = detect_eelarve_layout(worksheet)
        except ValueError:
            continue
        score = len(layout.month_columns)
        if score > best_score:
            best_score = score
            best_sheet = worksheet
    if best_score < 1:
        raise ValueError("No Eelarve sheet with month headers found")
    return best_sheet


def _row_is_total(worksheet: Worksheet, row: int, layout: EelarveLayout) -> bool:
    return is_total_label(cell_text(worksheet.cell(row, layout.konto_col).value)) or is_total_label(
        cell_text(worksheet.cell(row, layout.object_col).value)
    )


def _row_konto(worksheet: Worksheet, row: int, layout: EelarveLayout) -> str | None:
    if _row_is_total(worksheet, row, layout):
        return None
    return parse_konto(worksheet.cell(row, layout.konto_col).value) or parse_konto(
        worksheet.cell(row, layout.object_col).value
    )


def find_konto_blocks(worksheet: Worksheet, layout: EelarveLayout) -> list[KontoBlock]:
    blocks: list[KontoBlock] = []
    current: KontoBlock | None = None
    last = last_used_row(worksheet, layout)
    for row in range(layout.header_row + 1, last + 1):
        konto = _row_konto(worksheet, row, layout)
        if konto:
            if current is not None:
                blocks.append(current)
            current = KontoBlock(
                konto=konto,
                header_row=row,
                end_row=row,
                total_rows=[],
                labels={},
            )
            continue
        if current is None:
            continue
        if _row_is_total(worksheet, row, layout):
            current.total_rows.append(row)
            current.end_row = row
            continue
        object_label = cell_text(worksheet.cell(row, layout.object_col).value)
        konto_label = cell_text(worksheet.cell(row, layout.konto_col).value)
        if object_label or konto_label:
            current.end_row = row
            _index_object_row(current, worksheet, layout, row)
    if current is not None:
        blocks.append(current)
    return blocks


def _index_object_row(
    block: KontoBlock, worksheet: Worksheet, layout: EelarveLayout, row: int
) -> None:
    object_label = cell_text(worksheet.cell(row, layout.object_col).value)
    full = _row_search_text(worksheet, row, layout)
    for part in (object_label, full):
        key = _normalize_match_text(part)
        if key:
            block.labels[key] = row


def shift_blocks_after_insert(blocks: list[KontoBlock], insert_at: int, count: int = 1) -> None:
    for block in blocks:
        if block.header_row >= insert_at:
            block.header_row += count
        if block.end_row >= insert_at:
            block.end_row += count
        block.total_rows = [row + count if row >= insert_at else row for row in block.total_rows]
        block.labels = {
            key: (row + count if row >= insert_at else row) for key, row in block.labels.items()
        }


def first_blocks_by_konto(blocks: list[KontoBlock]) -> dict[str, KontoBlock]:
    index: dict[str, KontoBlock] = {}
    for block in blocks:
        index.setdefault(block.konto, block)
    return index


def _shift_row(row: int, insert_at: int, count: int) -> int:
    return row + count if row >= insert_at else row


def _adjust_sum_range(start_text: str, end_text: str, insert_at: int, count: int) -> str:
    start_match = CELL_RE.fullmatch(start_text)
    end_match = CELL_RE.fullmatch(end_text)
    if not start_match or not end_match:
        return f"{start_text}:{end_text}"
    start_col = start_match.group(2).upper()
    end_col = end_match.group(2).upper()
    start_row = int(start_match.group(4))
    end_row = int(end_match.group(4))
    if start_col == end_col:
        low, high = min(start_row, end_row), max(start_row, end_row)
        if low <= insert_at <= high + 1:
            new_low, new_high = low, high + count
        elif insert_at < low:
            new_low, new_high = low + count, high + count
        else:
            new_low, new_high = low, high
        if start_row <= end_row:
            start_row, end_row = new_low, new_high
        else:
            start_row, end_row = new_high, new_low
    else:
        start_row = _shift_row(start_row, insert_at, count)
        end_row = _shift_row(end_row, insert_at, count)
    start_out = f"{start_match.group(1)}{start_col}{start_match.group(3)}{start_row}"
    end_out = f"{end_match.group(1)}{end_col}{end_match.group(3)}{end_row}"
    return f"{start_out}:{end_out}"


def _adjust_sum_inner(inner: str, insert_at: int, count: int) -> str:
    pieces: list[str] = []
    cursor = 0
    for match in RANGE_RE.finditer(inner):
        pieces.append(inner[cursor : match.start()])
        pieces.append(_adjust_sum_range(match.group(1), match.group(2), insert_at, count))
        cursor = match.end()
    tail = inner[cursor:]

    def shift_cell(match: re.Match[str]) -> str:
        row = _shift_row(int(match.group(4)), insert_at, count)
        return f"{match.group(1)}{match.group(2).upper()}{match.group(3)}{row}"

    pieces.append(CELL_RE.sub(shift_cell, tail))
    return "".join(pieces)


def adjust_formula(formula: str, insert_at: int, count: int = 1) -> str:
    def replace_sum(match: re.Match[str]) -> str:
        return f"SUM({_adjust_sum_inner(match.group(1), insert_at, count)})"

    return SUM_CALL_RE.sub(replace_sum, formula)


def adjust_sum_formulas_after_insert(
    worksheet: Worksheet,
    insert_at: int,
    count: int = 1,
    min_row: int = 1,
    max_row: int | None = None,
) -> None:
    last = max_row if max_row is not None else worksheet.max_row
    first = max(1, min_row)
    for row in worksheet.iter_rows(min_row=first, max_row=last, max_col=worksheet.max_column):
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("=") and "SUM(" in value.upper():
                formula = value[1:]
                adjusted = adjust_formula(formula, insert_at, count)
                if adjusted != formula:
                    cell.value = "=" + adjusted


def create_blank_eelarve() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Eelarve"
    sheet.append(["Konto", "Object", *MONTHS, "Kokku"])
    sheet.freeze_panes = "A2"
    return workbook


def _month_col(layout: EelarveLayout, month: str) -> int:
    if month not in layout.month_columns:
        raise ValueError(f"Month column {month!r} not found in Eelarve headers")
    return layout.month_columns[month]


def _object_rows(block: KontoBlock) -> list[int]:
    skip = set(block.total_rows) | {block.header_row}
    return [row for row in range(block.header_row + 1, block.end_row + 1) if row not in skip]


def _row_search_text(worksheet: Worksheet, row: int, layout: EelarveLayout) -> str:
    return " ".join(
        part
        for part in (
            cell_text(worksheet.cell(row, layout.object_col).value),
            cell_text(worksheet.cell(row, layout.konto_col).value),
        )
        if part
    )


def find_object_row(
    worksheet: Worksheet,
    layout: EelarveLayout,
    block: KontoBlock,
    object_code: str,
    lookup: dict[str, str],
) -> int | None:
    description = lookup.get(object_code, "")
    for needle in (object_code, description):
        if not needle:
            continue
        row = block.labels.get(_normalize_match_text(needle))
        if row and row != block.header_row and row not in block.total_rows:
            return row
    for row in _object_rows(block):
        text = _row_search_text(worksheet, row, layout)
        if matches_label(text, object_code):
            _index_object_row(block, worksheet, layout, row)
            return row
        if description and matches_label(text, description):
            _index_object_row(block, worksheet, layout, row)
            return row
    return None


def _first_last_month_cols(layout: EelarveLayout) -> tuple[int, int]:
    columns = list(layout.month_columns.values())
    return min(columns), max(columns)


def write_row_month_total(worksheet: Worksheet, layout: EelarveLayout, row: int) -> None:
    if layout.total_col is None or not layout.month_columns:
        return
    first_col, last_col = _first_last_month_cols(layout)
    worksheet.cell(row, layout.total_col).value = (
        f"=SUM({get_column_letter(first_col)}{row}:{get_column_letter(last_col)}{row})"
    )


def ensure_block_formulas(worksheet: Worksheet, layout: EelarveLayout, block: KontoBlock) -> None:
    object_rows = _object_rows(block)
    if not object_rows:
        return
    first_obj, last_obj = min(object_rows), max(object_rows)
    for month_col in layout.month_columns.values():
        start = f"{get_column_letter(month_col)}{first_obj}"
        end = f"{get_column_letter(month_col)}{last_obj}"
        worksheet.cell(block.header_row, month_col).value = f"=SUM({start}:{end})"
    write_row_month_total(worksheet, layout, block.header_row)
    for row in object_rows:
        write_row_month_total(worksheet, layout, row)


def insert_object_row(
    worksheet: Worksheet,
    layout: EelarveLayout,
    block: KontoBlock,
    label: str,
    with_formulas: bool,
    blocks: list[KontoBlock] | None = None,
    last_row: int | None = None,
) -> int:
    insert_at = min(block.total_rows) if block.total_rows else block.end_row + 1
    used_last = last_row if last_row is not None else last_used_row(worksheet, layout)
    if insert_at <= used_last:
        worksheet.insert_rows(insert_at)
        if blocks is not None:
            shift_blocks_after_insert(blocks, insert_at)
        adjust_sum_formulas_after_insert(
            worksheet,
            insert_at,
            min_row=min(block.header_row, insert_at),
            max_row=max(used_last + 1, insert_at),
        )
    else:
        adjust_sum_formulas_after_insert(
            worksheet,
            insert_at,
            min_row=block.header_row,
            max_row=max(block.end_row, block.header_row),
        )
    worksheet.cell(insert_at, layout.object_col).value = label
    if with_formulas:
        write_row_month_total(worksheet, layout, insert_at)
    _index_object_row(block, worksheet, layout, insert_at)
    if insert_at > block.end_row:
        block.end_row = insert_at
    return insert_at


def append_konto_block(worksheet: Worksheet, layout: EelarveLayout, konto: str) -> int:
    row = last_used_row(worksheet) + 1
    worksheet.cell(row, layout.konto_col).value = konto
    return row


def apply_facts(
    worksheet: Worksheet,
    facts: list[SourceFact],
    month: str,
    lookup: dict[str, str],
    with_formulas: bool,
    on_written: Callable[[int], None] | None = None,
    should_cancel: Callable[[], bool] | None = None,
    cancel_every: int = CANCEL_CHECK_EVERY,
    layout: EelarveLayout | None = None,
    blocks: list[KontoBlock] | None = None,
) -> tuple[int, int, int]:
    written = 0
    rows_added = 0
    missing = 0
    seen_missing: set[str] = set()
    month_name = parse_month(month)
    layout = layout or detect_eelarve_layout(worksheet)
    if blocks is None:
        blocks = find_konto_blocks(worksheet, layout)
    by_konto = first_blocks_by_konto(blocks)
    month_col = _month_col(layout, month_name)
    last_row = last_used_row(worksheet, layout)

    for index, fact in enumerate(facts, start=1):
        if should_cancel is not None and index % max(1, cancel_every) == 0 and should_cancel():
            raise JobAborted(INTERRUPT_MESSAGE)
        block = by_konto.get(fact.konto)
        created_block = False
        if block is None:
            if fact.konto not in seen_missing:
                seen_missing.add(fact.konto)
                missing += 1
            last_row = max(last_row, last_used_row(worksheet, layout))
            header_row = last_row + 1
            worksheet.cell(header_row, layout.konto_col).value = fact.konto
            block = KontoBlock(
                konto=fact.konto,
                header_row=header_row,
                end_row=header_row,
                total_rows=[],
                labels={},
            )
            blocks.append(block)
            by_konto[fact.konto] = block
            last_row = header_row
            created_block = True

        had_objects = bool(_object_rows(block))
        object_row = find_object_row(worksheet, layout, block, fact.object_code, lookup)
        if object_row is None:
            label = lookup.get(fact.object_code) or fact.object_code
            object_row = insert_object_row(
                worksheet,
                layout,
                block,
                label,
                with_formulas,
                blocks=blocks,
                last_row=last_row,
            )
            last_row = max(last_row + 1, object_row, block.end_row)
            rows_added += 1
            if with_formulas and not had_objects:
                ensure_block_formulas(worksheet, layout, block)
        elif created_block and with_formulas:
            ensure_block_formulas(worksheet, layout, block)

        worksheet.cell(object_row, month_col).value = fact.amount
        written += 1
        if on_written is not None:
            on_written(1)
    return written, rows_added, missing


def fill_eelarve(
    *,
    eelarve_path: Path | None,
    sources: list[tuple[Path, str]],
    objects_path: Path | None = None,
    output_path: Path | None = None,
    year: int | None = None,
    root: Path | None = None,
    reporter: ProgressReporter | None = None,
    use_gui: bool = False,
    handle_locks: bool = False,
    should_cancel: Callable[[], bool] | None = None,
) -> FillStats:
    if not sources:
        raise ValueError("No source workbooks provided")
    reporter = reporter or ProgressReporter(enabled=False)
    created_new = eelarve_path is None
    output = Path(output_path) if output_path is not None else None
    workbook = None
    try:
        if created_new:
            workbook = create_blank_eelarve()
            worksheet = workbook.active
            if output is None:
                output = default_new_eelarve_path(root, year)
            output.parent.mkdir(parents=True, exist_ok=True)
            reporter.stage(5, f"Создан новый Eelarve: {_resolved_path_text(output)}")
        else:
            reporter.stage(5, f"Eelarve выбран: {_resolved_path_text(Path(eelarve_path))}")
            status, workbook = run_with_lock_handling(
                lambda: load_workbook(eelarve_path),
                Path(eelarve_path),
                kind="eelarve",
                handle_locks=handle_locks,
                use_gui=use_gui,
            )
            if status != "ok" or workbook is None:
                raise JobAborted(
                    "Работа прервана: не удалось открыть Eelarve.\n"
                    f"{_resolved_path_text(Path(eelarve_path))}"
                )
            worksheet = pick_eelarve_sheet(workbook)
            if output is None:
                output = Path(eelarve_path)
            output.parent.mkdir(parents=True, exist_ok=True)

        def checkpoint(label: str) -> None:
            reporter.stage(max(reporter.percent, 6), f"Сохранение ({label}): {_resolved_path_text(output)}")
            run_with_lock_handling(
                lambda: workbook.save(output),
                output,
                kind="output",
                handle_locks=handle_locks,
                use_gui=use_gui,
            )

        if created_new:
            checkpoint("заголовок")

        for source_path, month in sources:
            if should_cancel is not None and should_cancel():
                raise JobAborted(INTERRUPT_MESSAGE)
            reporter.stage(5, f"Источник принят: {Path(source_path).name} ({parse_month(month)})")

        if objects_path is not None:
            status, lookup = run_with_lock_handling(
                lambda: load_object_lookup(objects_path),
                Path(objects_path),
                kind="objects",
                handle_locks=handle_locks,
                use_gui=use_gui,
            )
            if status == "skip" or lookup is None:
                lookup = {}
                reporter.stage(8, f"Справочник объектов пропущен: {Path(objects_path).name}")
            else:
                reporter.stage(8, f"Справочник объектов загружен: {Path(objects_path).name}")
        else:
            lookup = {}
            reporter.stage(8, "Справочник объектов пропущен")

        parsed: list[tuple[Path, str, list[SourceFact]]] = []
        source_count = len(sources)
        for index, (source_path, month) in enumerate(sources, start=1):
            if should_cancel is not None and should_cancel():
                raise JobAborted(INTERRUPT_MESSAGE)
            source_path = Path(source_path)
            parse_pct = 10 + (20 * (index - 1) // max(source_count, 1))
            reporter.stage(parse_pct, f"Разбор источника {index} из {source_count}: {source_path.name}")
            status, facts = run_with_lock_handling(
                lambda path=source_path: parse_kasumiaruanne(path),
                source_path,
                kind="source",
                handle_locks=handle_locks,
                use_gui=use_gui,
            )
            if status == "skip" or facts is None:
                reporter.stage(
                    reporter.percent if reporter.percent >= 0 else parse_pct,
                    f"Источник пропущен: {source_path.name}",
                )
                continue
            reporter.stage(parse_pct, f"Разобрано сумм: {len(facts)} ({source_path.name})")
            parsed.append((source_path, month, facts))

        if not parsed:
            raise JobAborted("Все исходные файлы пропущены.\nAll source files were skipped.")

        total_facts = sum(len(facts) for _, _, facts in parsed)
        reporter.begin_fill(total_facts)
        layout = detect_eelarve_layout(worksheet)
        blocks = find_konto_blocks(worksheet, layout)

        written = 0
        rows_added = 0
        missing = 0
        for source_path, month, facts in parsed:
            if should_cancel is not None and should_cancel():
                raise JobAborted(INTERRUPT_MESSAGE)
            added_written, added_rows, added_missing = apply_facts(
                worksheet,
                facts,
                month,
                lookup,
                with_formulas=created_new,
                on_written=reporter.add_filled if reporter.enabled else None,
                should_cancel=should_cancel,
                layout=layout,
                blocks=blocks,
            )
            written += added_written
            rows_added += added_rows
            if not created_new:
                missing += added_missing
            checkpoint(source_path.name)

        reporter.stage(95, f"Сохранение: {_resolved_path_text(output)}")
        checkpoint("итог")
        reporter.stage(100, "Готово")
        return FillStats(
            written=written,
            rows_added=rows_added,
            missing_konto_blocks=missing,
            output_path=output,
        )
    except KeyboardInterrupt as exc:
        raise JobAborted(INTERRUPT_MESSAGE) from exc


def prompt_text_paths() -> list[Path]:
    print("Enter source Excel .xlsx files one per line. Press Enter on an empty line when done.")
    files: list[Path] = []
    while True:
        try:
            value = input("Source file path: ").strip().strip('"')
        except EOFError:
            break
        if not value:
            break
        files.append(Path(value))
    return files


def prompt_interactive(args: argparse.Namespace) -> argparse.Namespace:
    input_dir, output_dir = ensure_work_dirs()
    year = args.year if getattr(args, "year", None) else date.today().year
    default_new = output_dir / f"Eelarve_{year}.xlsx"
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, simpledialog

        root = tk.Tk()
        root.withdraw()
        messagebox.showinfo(
            "Excellent Books Eelarve",
            "Vali olemasolev Eelarve töövihik või tühista, et luua uus.\n"
            "Select an existing Eelarve workbook, or Cancel to create a new one.",
        )
        selected_eelarve = filedialog.askopenfilename(
            title="Vali olemasolev Eelarve (Cancel = loo uus / Cancel = create new)",
            initialdir=str(input_dir),
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
        )
        selected_files = filedialog.askopenfilenames(
            title="Vali igakuised Kasumiaruanne .xlsx failid",
            initialdir=str(input_dir),
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
        )
        if not selected_files:
            root.destroy()
            raise SystemExit("Исходные файлы не выбраны.\nNo source files selected.")

        months: list[str] = []
        for path in selected_files:
            month_text = simpledialog.askstring(
                "Kuu / Month",
                f"Month for {Path(path).name} (Jaanuar..Detsember or 1-12):",
                initialvalue="Jaanuar",
            )
            if not month_text:
                root.destroy()
                raise SystemExit("Kuu ei ole valitud.\nNo month selected.")
            months.append(parse_month(month_text))

        object_lookup = filedialog.askopenfilename(
            title="Optional: object lookup CSV/XLSX, or Cancel to skip",
            initialdir=str(input_dir),
            filetypes=[
                ("Object lookup", "*.csv *.xlsx"),
                ("CSV files", "*.csv"),
                ("Excel files", "*.xlsx"),
                ("All files", "*.*"),
            ],
        )
        root.destroy()

        args.eelarve = Path(selected_eelarve) if selected_eelarve else None
        args.sources = [Path(path) for path in selected_files]
        args.months = months
        args.objects = Path(object_lookup) if object_lookup else None
        if args.output is None:
            args.output = args.eelarve if args.eelarve is not None else default_new
        if args.eelarve is not None:
            print(f"Eelarve выбран: {args.eelarve}", flush=True)
        else:
            print(f"Создан новый Eelarve: {default_new}", flush=True)
        for path, month in zip(args.sources, args.months, strict=True):
            print(f"Источник принят: {path.name} ({month})", flush=True)
        if args.objects is not None:
            print(f"Справочник объектов: {args.objects.name}", flush=True)
        else:
            print("Справочник объектов пропущен", flush=True)
        return args
    except ModuleNotFoundError:
        print("tkinter is not available; falling back to console prompts.")
    except tk.TclError:  # type: ignore[name-defined]  # pragma: no cover - GUI-specific
        print("GUI is not available; falling back to console prompts.")

    try:
        eelarve_text = input(
            "Eelarve workbook path (Enter = create new): "
        ).strip().strip('"')
    except EOFError as exc:
        raise SystemExit("Ввод отменен.\nInput cancelled.") from exc
    args.eelarve = Path(eelarve_text) if eelarve_text else None
    selected_files = prompt_text_paths()
    if not selected_files:
        raise SystemExit("Исходные файлы не выбраны.\nNo source files selected.")
    months = []
    for path in selected_files:
        try:
            month_text = input(
                f"Month for {path.name} (Jaanuar..Detsember or 1-12): "
            ).strip()
        except EOFError as exc:
            raise SystemExit("Ввод отменен.\nInput cancelled.") from exc
        months.append(parse_month(month_text or "Jaanuar"))
    try:
        object_lookup_text = input(
            "Optional object lookup CSV/XLSX path (press Enter to skip): "
        ).strip().strip('"')
    except EOFError as exc:
        raise SystemExit("Ввод отменен.\nInput cancelled.") from exc
    args.sources = selected_files
    args.months = months
    args.objects = Path(object_lookup_text) if object_lookup_text else None
    if args.output is None:
        args.output = args.eelarve if args.eelarve is not None else default_new
    if args.eelarve is not None:
        print(f"Eelarve выбран: {args.eelarve}", flush=True)
    else:
        print(f"Создан новый Eelarve: {default_new}", flush=True)
    for path, month in zip(args.sources, args.months, strict=True):
        print(f"Источник принят: {path.name} ({month})", flush=True)
    if args.objects is not None:
        print(f"Справочник объектов: {args.objects.name}", flush=True)
    else:
        print("Справочник объектов пропущен", flush=True)
    return args


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Fill an Eelarve workbook from monthly Kasumiaruanne Excel exports."
    )
    parser.add_argument(
        "--interactive",
        action="store_true",
        help="Ask for Eelarve, source files, months, and optional object lookup.",
    )
    parser.add_argument("--eelarve", type=Path, help="Existing Eelarve workbook. Omit to create a new one.")
    parser.add_argument(
        "--source",
        dest="sources",
        action="append",
        type=Path,
        default=[],
        help="Kasumiaruanne .xlsx file. Repeatable; pair with --month.",
    )
    parser.add_argument(
        "--month",
        dest="months",
        action="append",
        default=[],
        help="Month for the corresponding --source (Jaanuar..Detsember or 1-12).",
    )
    parser.add_argument(
        "--objects",
        type=Path,
        help="Optional object lookup .csv/.xlsx with object_code and object_name.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="Output path. Defaults to the selected Eelarve, or output/Eelarve_YYYY.xlsx.",
    )
    parser.add_argument(
        "--year",
        type=int,
        help="Year used in the default new filename Eelarve_YYYY.xlsx.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    interactive = args.interactive
    reporter = ProgressReporter(enabled=True)
    try:
        reporter.stage(0, "Запуск…")
        if interactive:
            args = prompt_interactive(args)
        if not args.sources:
            notify_user(
                "Excellent Books",
                "Исходные файлы не выбраны.\nNo source files selected.",
                error=True,
                use_gui=interactive,
            )
            return 2
        if len(args.sources) != len(args.months):
            raise ValueError("Each --source needs a matching --month")
        stats = fill_eelarve(
            eelarve_path=args.eelarve,
            sources=list(zip(args.sources, args.months, strict=True)),
            objects_path=args.objects,
            output_path=args.output,
            year=args.year,
            reporter=reporter,
            use_gui=interactive,
            handle_locks=True,
        )
        resolved = stats.output_path.expanduser().resolve()
        message = (
            "Eelarve заполнен.\n\n"
            f"Записано: {stats.written}\n"
            f"Добавлено строк: {stats.rows_added}\n"
            f"Нет блока konto: {stats.missing_konto_blocks}\n"
            f"{resolved}"
        )
        notify_user("Excellent Books", message, use_gui=interactive)
        return 0
    except KeyboardInterrupt:
        notify_user("Excellent Books", INTERRUPT_MESSAGE, error=True, use_gui=False)
        return 130
    except JobAborted as exc:
        interrupted = "Ctrl+C" in str(exc)
        notify_user(
            "Excellent Books",
            str(exc),
            error=True,
            use_gui=interactive and not interrupted,
        )
        return 130 if interrupted else 1
    except SystemExit as exc:
        code = exc.code
        if code is None or code == 0:
            raise
        if isinstance(code, str):
            notify_user("Excellent Books", code, error=True, use_gui=interactive)
            return 1
        return int(code)
    except Exception as exc:  # noqa: BLE001 - interactive clients need a clear message
        notify_user(
            "Excellent Books",
            f"Ошибка / Error:\n{exc}",
            error=True,
            use_gui=interactive,
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
