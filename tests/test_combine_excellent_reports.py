import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.combine_excellent_reports import main


class CombineExcellentReportsTests(unittest.TestCase):
    def test_unpivots_object_columns_to_csv(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "Kinnisvara.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Report"
            sheet.append(["Account", "Name", "HK_NOMME", "HK_KESKLINN"])
            sheet.append(["4000", "Revenue", 100, 200])
            sheet.append(["5000", "Cost", None, 50])
            workbook.save(source)

            objects = root / "objects.csv"
            objects.write_text(
                "object_code,object_name\nHK_NOMME,Nomme LOV\nHK_KESKLINN,Kesklinn LOV\n",
                encoding="utf-8",
            )

            output = root / "combined.csv"
            exit_code = main(
                [
                    "--file",
                    str(source),
                    "--objects",
                    str(objects),
                    "--output",
                    str(output),
                ]
            )

            self.assertEqual(exit_code, 0)
            with output.open("r", encoding="utf-8-sig") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(len(rows), 3)
            self.assertEqual(rows[0]["source_file"], "Kinnisvara.xlsx")
            self.assertEqual(rows[0]["object_code"], "HK_NOMME")
            self.assertEqual(rows[0]["object_name"], "Nomme LOV")
            self.assertEqual(rows[0]["amount"], "100")
            self.assertEqual(rows[0]["Account"], "4000")

    def test_writes_xlsx_with_summary_and_totals(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "Eelarve - 2027 var 2026.09.10.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Budget"
            sheet.append(["skip", "skip", "skip"])
            sheet.append(["Account", "Name", "HK_NOMME"])
            sheet.append(["4000", "Revenue", 125.5])
            workbook.save(source)

            output = root / "combined.xlsx"
            exit_code = main(
                [
                    "--file",
                    str(source),
                    "--output",
                    str(output),
                ]
            )

            self.assertEqual(exit_code, 0)
            result = load_workbook(output, read_only=True, data_only=True)
            self.assertEqual(result.sheetnames, ["combined", "summary", "totals_by_object"])
            combined = result["combined"]
            headers = [cell.value for cell in next(combined.iter_rows(max_row=1))]
            values = [cell.value for cell in next(combined.iter_rows(min_row=2, max_row=2))]
            row = dict(zip(headers, values, strict=False))
            self.assertEqual(row["object_code"], "HK_NOMME")
            self.assertEqual(row["amount"], 125.5)
            self.assertEqual(row["source_row"], 3)


if __name__ == "__main__":
    unittest.main()
