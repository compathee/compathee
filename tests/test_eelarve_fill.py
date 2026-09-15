import importlib.util
import io
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
CLIENT_SCRIPT = ROOT / "excellent" / "scripts" / "eelarve_fill.py"


def load_fill():
    spec = importlib.util.spec_from_file_location("eelarve_fill", CLIENT_SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def write_kasumiaruanne(path: Path, rows: list[list[object]], objects: list[str] | None = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kasumiaruanne"
    object_codes = objects or ["HK_A", "HK_B"]
    sheet.append(["Konto", "Nimetus", "Periood", *object_codes])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def write_existing_eelarve(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PL"
    sheet.append(["Konto", "Objekt", "Jaanuar", "Veebruar", "Kokku"])
    sheet.append(["3241 Tulud", None, None, None, None])
    sheet.append([None, "HK_A", None, None, None])
    sheet.append(["KOKKU", None, 999, None, None])
    sheet.append(["6000 Üldkulud", "no-objects", 777, None, None])
    sheet.append(["5000 Kulud", None, None, None, None])
    sheet.append([None, "HK_A", None, None, None])
    workbook.save(path)


class EelarveFillHelperTests(unittest.TestCase):
    def test_parse_amount_supports_spaced_decimal_comma(self):
        fill = load_fill()
        self.assertEqual(fill.parse_amount("2 848,00"), 2848.0)
        self.assertEqual(fill.parse_amount("2\u00a0848,00"), 2848.0)
        self.assertEqual(fill.parse_amount(100), 100.0)
        self.assertEqual(fill.parse_amount(125.5), 125.5)
        self.assertIsNone(fill.parse_amount(None))
        self.assertIsNone(fill.parse_amount(""))
        self.assertIsNone(fill.parse_amount("  "))

    def test_parse_konto_extracts_account_number(self):
        fill = load_fill()
        self.assertEqual(fill.parse_konto("3241"), "3241")
        self.assertEqual(fill.parse_konto("3241 Tulu"), "3241")
        self.assertEqual(fill.parse_konto(3241.0), "3241")
        self.assertIsNone(fill.parse_konto("Budget 2026"))
        self.assertIsNone(fill.parse_konto("Tulud"))

    def test_is_total_label_detects_kokku_rows(self):
        fill = load_fill()
        self.assertTrue(fill.is_total_label("Kokku tulu"))
        self.assertTrue(fill.is_total_label("KOKKU"))
        self.assertFalse(fill.is_total_label("HK_NOMME"))
        self.assertFalse(fill.is_total_label("3241 Müügitulu"))

    def test_default_new_eelarve_path_uses_year(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = fill.default_new_eelarve_path(root, year=2026)
            self.assertEqual(path, root / "output" / "Eelarve_2026.xlsx")

    def test_progress_percent_maps_fill_stage(self):
        fill = load_fill()
        self.assertEqual(fill.progress_percent(0, 10, start=30, end=90), 30)
        self.assertEqual(fill.progress_percent(5, 10, start=30, end=90), 60)
        self.assertEqual(fill.progress_percent(10, 10, start=30, end=90), 90)
        self.assertEqual(fill.progress_percent(0, 0, start=30, end=90), 90)

    def test_is_lock_error_detects_permission_and_sharing(self):
        fill = load_fill()
        self.assertTrue(fill.is_lock_error(PermissionError("denied")))
        self.assertTrue(fill.is_lock_error(OSError(13, "Permission denied")))
        busy = OSError("The process cannot access the file because it is being used by another process")
        busy.winerror = 32
        self.assertTrue(fill.is_lock_error(busy))
        self.assertFalse(fill.is_lock_error(ValueError("bad month")))


class EelarveFillEngineTests(unittest.TestCase):
    def test_parse_kasumiaruanne_skips_period_and_kokku(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "jaanuar.xlsx"
            write_kasumiaruanne(
                source,
                [
                    ["3241", "Tulu", "2026-01", "2 848,00", 100],
                    ["Kokku tulu", "", "2026-01", 9999, 8888],
                ],
                objects=["HK_NOMME", "KÜ_TEST"],
            )
            facts = fill.parse_kasumiaruanne(source)
            pairs = {(fact.konto, fact.object_code, fact.amount) for fact in facts}
            self.assertEqual(
                pairs,
                {("3241", "HK_NOMME", 2848.0), ("3241", "KÜ_TEST", 100.0)},
            )

    def test_block_match_writes_only_inside_matching_konto(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            eelarve = root / "eelarve.xlsx"
            output = root / "out.xlsx"
            write_kasumiaruanne(source, [["3241", "Tulu", "2026-01", 100, None]])
            write_existing_eelarve(eelarve)
            stats = fill.fill_eelarve(
                eelarve_path=eelarve,
                sources=[(source, "Jaanuar")],
                output_path=output,
            )
            self.assertGreaterEqual(stats.written, 1)
            result = load_workbook(output)
            sheet = result["PL"]
            self.assertEqual(sheet["C3"].value, 100)
            self.assertNotEqual(sheet["C7"].value, 100)
            self.assertIsNone(sheet["C7"].value)
            self.assertEqual(sheet["C5"].value, 777)

    def test_csv_description_match(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            eelarve = root / "eelarve.xlsx"
            objects = root / "objects.csv"
            output = root / "out.xlsx"
            write_kasumiaruanne(
                source,
                [["3241", "Tulu", "2026-01", 50]],
                objects=["HK_NOMME"],
            )
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "PL"
            sheet.append(["Konto", "Objekt", "Jaanuar"])
            sheet.append(["3241 Tulud", None, None])
            sheet.append([None, "Nomme LOV", None])
            workbook.save(eelarve)
            objects.write_text(
                "object_code,object_name\nHK_NOMME,Nomme LOV\n",
                encoding="utf-8",
            )
            fill.fill_eelarve(
                eelarve_path=eelarve,
                sources=[(source, "Jaanuar")],
                objects_path=objects,
                output_path=output,
            )
            sheet = load_workbook(output)["PL"]
            self.assertEqual(sheet["C3"].value, 50)

    def test_append_on_miss_and_skip_unrelated_section(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            eelarve = root / "eelarve.xlsx"
            output = root / "out.xlsx"
            write_kasumiaruanne(source, [["3241", "Tulu", "2026-01", None, 40]])
            write_existing_eelarve(eelarve)
            stats = fill.fill_eelarve(
                eelarve_path=eelarve,
                sources=[(source, "Jaanuar")],
                output_path=output,
            )
            self.assertEqual(stats.rows_added, 1)
            sheet = load_workbook(output)["PL"]
            self.assertEqual(sheet["B4"].value, "HK_B")
            self.assertEqual(sheet["C4"].value, 40)
            self.assertEqual(sheet["A5"].value, "KOKKU")
            self.assertEqual(sheet["C5"].value, 999)
            self.assertEqual(sheet["C6"].value, 777)

    def test_plain_konto_row_gets_object_children(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            eelarve = root / "eelarve.xlsx"
            output = root / "out.xlsx"
            write_kasumiaruanne(source, [["3241", "Tulu", "2026-01", 12, None]])
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Konto", "Objekt", "Jaanuar"])
            sheet.append(["3241 Tulud", None, None])
            sheet.append(["5000 Kulud", "plain", 42])
            workbook.save(eelarve)
            fill.fill_eelarve(
                eelarve_path=eelarve,
                sources=[(source, "Jaanuar")],
                output_path=output,
            )
            sheet = load_workbook(output).active
            self.assertEqual(sheet["B3"].value, "HK_A")
            self.assertEqual(sheet["C3"].value, 12)
            self.assertEqual(sheet["A4"].value, "5000 Kulud")
            self.assertEqual(sheet["C4"].value, 42)

    def test_does_not_overwrite_kokku_values(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            eelarve = root / "eelarve.xlsx"
            output = root / "out.xlsx"
            write_kasumiaruanne(
                source,
                [["3241", "KOKKU should not match", "2026-01", 1, None]],
                objects=["KOKKU"],
            )
            write_existing_eelarve(eelarve)
            fill.fill_eelarve(
                eelarve_path=eelarve,
                sources=[(source, "Jaanuar")],
                output_path=output,
            )
            sheet = load_workbook(output)["PL"]
            kokku_values = [
                sheet.cell(row, 3).value
                for row in range(1, sheet.max_row + 1)
                if fill.is_total_label(str(sheet.cell(row, 1).value or ""))
            ]
            self.assertIn(999, kokku_values)

    def test_create_new_eelarve_naming_and_sum_formulas(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            write_kasumiaruanne(source, [["3241", "Tulu", "2026-01", 10, 5]])
            stats = fill.fill_eelarve(
                eelarve_path=None,
                sources=[(source, "Jaanuar")],
                year=2026,
                root=root,
            )
            self.assertEqual(stats.output_path, root / "output" / "Eelarve_2026.xlsx")
            self.assertTrue(stats.output_path.is_file())
            self.assertEqual(stats.missing_konto_blocks, 0)
            sheet = load_workbook(stats.output_path)["Eelarve"]
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            self.assertEqual(headers[0], "Konto")
            self.assertEqual(headers[1], "Object")
            self.assertIn("Jaanuar", headers)
            self.assertIn("Detsember", headers)
            self.assertIn("Kokku", headers)
            self.assertEqual(sheet["A2"].value, "3241")
            formulas = [
                str(sheet.cell(row, column).value)
                for row in range(1, sheet.max_row + 1)
                for column in range(1, sheet.max_column + 1)
                if isinstance(sheet.cell(row, column).value, str)
                and str(sheet.cell(row, column).value).startswith("=")
            ]
            self.assertTrue(any("SUM(" in formula.upper() for formula in formulas))

    def test_inserting_object_expands_neighboring_sum(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            eelarve = root / "eelarve.xlsx"
            output = root / "out.xlsx"
            write_kasumiaruanne(source, [["3241", "Tulu", "2026-01", None, 7]])
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Konto", "Objekt", "Jaanuar"])
            sheet.append(["3241", None, "=SUM(C3:C3)"])
            sheet.append([None, "HK_A", 10])
            workbook.save(eelarve)
            fill.fill_eelarve(
                eelarve_path=eelarve,
                sources=[(source, "Jaanuar")],
                output_path=output,
            )
            sheet = load_workbook(output).active
            self.assertEqual(sheet["C2"].value, "=SUM(C3:C4)")
            self.assertEqual(sheet["B4"].value, "HK_B")
            self.assertEqual(sheet["C4"].value, 7)

    def test_missing_konto_appends_block_and_counts(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            eelarve = root / "eelarve.xlsx"
            output = root / "out.xlsx"
            write_kasumiaruanne(source, [["3241", "Tulu", "2026-01", 3, None]])
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Konto", "Objekt", "Jaanuar"])
            sheet.append(["5000 Kulud", "plain", 42])
            workbook.save(eelarve)
            stats = fill.fill_eelarve(
                eelarve_path=eelarve,
                sources=[(source, "Jaanuar")],
                output_path=output,
            )
            self.assertEqual(stats.missing_konto_blocks, 1)
            sheet = load_workbook(output).active
            self.assertEqual(sheet["C2"].value, 42)
            self.assertEqual(sheet["A3"].value, "3241")
            self.assertEqual(sheet["B4"].value, "HK_A")
            self.assertEqual(sheet["C4"].value, 3)

    def test_main_prints_stats_and_absolute_path(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            output = root / "nested" / "Eelarve_2026.xlsx"
            write_kasumiaruanne(source, [["3241", "Tulu", "2026-01", 1, None]])
            buf = io.StringIO()
            with redirect_stdout(buf):
                code = fill.main(
                    [
                        "--source",
                        str(source),
                        "--month",
                        "Jaanuar",
                        "--output",
                        str(output),
                        "--year",
                        "2026",
                    ]
                )
            self.assertEqual(code, 0)
            text = buf.getvalue()
            self.assertIn("Eelarve заполнен", text)
            self.assertIn("Записано:", text)
            self.assertIn(str(output.resolve()), text)
            self.assertRegex(text, r"\[\s*\d+%\]")
            self.assertIn("Заполнение Eelarve", text)
            self.assertNotIn("HK_A", text.split("Eelarve заполнен")[0])

    def test_interactive_cancel_eelarve_creates_new_file(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "input").mkdir()
            (root / "output").mkdir()
            source = root / "src.xlsx"
            write_kasumiaruanne(source, [["3241", "Tulu", "2026-01", 8, None]])
            stdin = io.StringIO(f"\n{source}\n\nJaanuar\n\n")
            stdout = io.StringIO()
            old_stdin = sys.stdin
            sys.stdin = stdin
            try:
                with patch.dict(sys.modules, {"tkinter": None}):
                    with patch.object(
                        fill, "ensure_work_dirs", return_value=(root / "input", root / "output")
                    ):
                        with redirect_stdout(stdout):
                            code = fill.main(["--interactive", "--year", "2026"])
            finally:
                sys.stdin = old_stdin
            self.assertEqual(code, 0)
            created = root / "output" / "Eelarve_2026.xlsx"
            self.assertTrue(created.is_file())
            self.assertEqual(load_workbook(created)["Eelarve"]["C3"].value, 8)

    def test_interactive_cancel_sources_returns_nonzero(self):
        fill = load_fill()
        stdin = io.StringIO("\n\n")
        stderr = io.StringIO()
        old_stdin = sys.stdin
        sys.stdin = stdin
        try:
            with patch.dict(sys.modules, {"tkinter": None}):
                with redirect_stdout(io.StringIO()), redirect_stderr(stderr):
                    code = fill.main(["--interactive"])
        finally:
            sys.stdin = old_stdin
        self.assertNotEqual(code, 0)
        self.assertIn("Исходные файлы не выбраны", stderr.getvalue())


class EelarveFillProgressAndLockTests(unittest.TestCase):
    def test_locked_source_abort_returns_nonzero(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            output = root / "out.xlsx"
            write_kasumiaruanne(source, [["3241", "Tulu", "2026-01", 1, None]])
            stdin = io.StringIO("2\n")
            stdout = io.StringIO()
            stderr = io.StringIO()
            old_stdin = sys.stdin
            sys.stdin = stdin
            try:
                with patch.object(fill, "parse_kasumiaruanne", side_effect=PermissionError("busy")):
                    with redirect_stdout(stdout), redirect_stderr(stderr):
                        code = fill.main(
                            [
                                "--source",
                                str(source),
                                "--month",
                                "Jaanuar",
                                "--output",
                                str(output),
                            ]
                        )
            finally:
                sys.stdin = old_stdin
            self.assertNotEqual(code, 0)
            combined = stdout.getvalue() + stderr.getvalue()
            self.assertIn("src.xlsx", combined)
            self.assertRegex(combined, r"(?i)занят|busy|locked")
            self.assertFalse(output.exists())

    def test_locked_source_continue_skips_and_proceeds(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            locked = root / "locked.xlsx"
            ok_source = root / "ok.xlsx"
            output = root / "out.xlsx"
            write_kasumiaruanne(locked, [["3241", "Tulu", "2026-01", 9, None]])
            write_kasumiaruanne(ok_source, [["3241", "Tulu", "2026-01", 4, None]])
            real_parse = fill.parse_kasumiaruanne

            def parse_maybe_locked(path):
                if Path(path).name == "locked.xlsx":
                    raise PermissionError("busy")
                return real_parse(path)

            stdin = io.StringIO("1\n")
            stdout = io.StringIO()
            old_stdin = sys.stdin
            sys.stdin = stdin
            try:
                with patch.object(fill, "parse_kasumiaruanne", side_effect=parse_maybe_locked):
                    with redirect_stdout(stdout), redirect_stderr(io.StringIO()):
                        code = fill.main(
                            [
                                "--source",
                                str(locked),
                                "--month",
                                "Jaanuar",
                                "--source",
                                str(ok_source),
                                "--month",
                                "Jaanuar",
                                "--output",
                                str(output),
                            ]
                        )
            finally:
                sys.stdin = old_stdin
            self.assertEqual(code, 0)
            self.assertTrue(output.exists())
            sheet = load_workbook(output)["Eelarve"]
            self.assertEqual(sheet["C3"].value, 4)
            self.assertIn("locked.xlsx", stdout.getvalue() + "")

    def test_locked_output_abort_does_not_claim_success(self):
        fill = load_fill()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "src.xlsx"
            output = root / "out.xlsx"
            write_kasumiaruanne(source, [["3241", "Tulu", "2026-01", 1, None]])
            stdin = io.StringIO("2\n")
            stdout = io.StringIO()
            stderr = io.StringIO()
            old_stdin = sys.stdin
            sys.stdin = stdin
            try:
                with patch.object(fill.Workbook, "save", side_effect=PermissionError("busy")):
                    with redirect_stdout(stdout), redirect_stderr(stderr):
                        code = fill.main(
                            [
                                "--source",
                                str(source),
                                "--month",
                                "Jaanuar",
                                "--output",
                                str(output),
                            ]
                        )
            finally:
                sys.stdin = old_stdin
            self.assertNotEqual(code, 0)
            combined = stdout.getvalue() + stderr.getvalue()
            self.assertIn("out.xlsx", combined)
            self.assertNotIn("Eelarve заполнен", combined)


if __name__ == "__main__":
    unittest.main()
